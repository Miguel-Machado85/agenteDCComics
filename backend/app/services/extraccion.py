"""
Extraccion de texto de archivos subidos (POST /documents).

Convierte los bytes de un archivo en una lista de `Seccion`, que luego recibe
`ingestar()`. Se parte por estructura para que cada chunk conserve contexto:
- PDF: una seccion por pagina;
- MD y DOCX: una seccion por titulo (el nombre de la seccion es el titulo);
- CSV: una seccion por fila, como texto "columna: valor";
- TXT: una sola seccion.
Las secciones cortas las vuelve a juntar el chunking de `ingestion.py`.
"""

import csv
import io
import re
from pathlib import Path

from app.services.ingestion import Seccion

FORMATOS = (".pdf", ".txt", ".md", ".csv", ".docx")


class FormatoNoSoportado(Exception):
    pass


class SinTexto(Exception):
    """El archivo no se pudo leer o no tiene texto extraible."""


def _decodificar(datos: bytes) -> str:
    try:
        return datos.decode("utf-8-sig")
    except UnicodeDecodeError:
        return datos.decode("latin-1")


def _desde_txt(datos: bytes) -> list[Seccion]:
    return [Seccion("Texto", _decodificar(datos))]


def _desde_md(datos: bytes) -> list[Seccion]:
    secciones: list[Seccion] = []
    nombre, lineas = "Introduccion", []
    for linea in _decodificar(datos).splitlines():
        titulo = re.match(r"^#{1,6}\s+(.+?)\s*#*$", linea)
        if titulo:
            secciones.append(Seccion(nombre, "\n".join(lineas)))
            nombre, lineas = titulo.group(1), []
        else:
            lineas.append(linea)
    secciones.append(Seccion(nombre, "\n".join(lineas)))
    return secciones


def _desde_csv(datos: bytes) -> list[Seccion]:
    texto = _decodificar(datos)
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = list(csv.reader(io.StringIO(texto), dialecto))
    if len(filas) < 2:
        return [Seccion("Texto", texto)]
    columnas = [c.strip() for c in filas[0]]
    secciones = []
    for n, fila in enumerate(filas[1:], 1):
        pares = [f"{c}: {v.strip()}" for c, v in zip(columnas, fila) if v.strip()]
        secciones.append(Seccion(f"Fila {n}", "\n".join(pares)))
    return secciones


def _desde_pdf(datos: bytes) -> list[Seccion]:
    from pypdf import PdfReader

    try:
        lector = PdfReader(io.BytesIO(datos))
        return [
            Seccion(f"Pagina {n}", pagina.extract_text() or "")
            for n, pagina in enumerate(lector.pages, 1)
        ]
    except Exception as exc:   # pypdf lanza varios tipos segun el dano
        raise SinTexto("No se pudo leer el PDF (esta danado o protegido).") from exc


def _desde_docx(datos: bytes) -> list[Seccion]:
    from docx import Document

    try:
        documento = Document(io.BytesIO(datos))
    except Exception as exc:
        raise SinTexto("No se pudo leer el DOCX (esta danado).") from exc
    secciones: list[Seccion] = []
    nombre, parrafos = "Introduccion", []
    for p in documento.paragraphs:
        estilo = (p.style.name or "") if p.style is not None else ""
        if estilo.startswith(("Heading", "Title")) and p.text.strip():
            secciones.append(Seccion(nombre, "\n".join(parrafos)))
            nombre, parrafos = p.text.strip(), []
        else:
            parrafos.append(p.text)
    secciones.append(Seccion(nombre, "\n".join(parrafos)))
    for tabla in documento.tables:
        filas = [" | ".join(c.text.strip() for c in f.cells) for f in tabla.rows]
        secciones.append(Seccion("Tabla", "\n".join(filas)))
    return secciones


EXTRACTORES = {
    ".pdf": _desde_pdf,
    ".txt": _desde_txt,
    ".md": _desde_md,
    ".csv": _desde_csv,
    ".docx": _desde_docx,
}


def extraer(nombre_archivo: str, datos: bytes) -> list[Seccion]:
    """Secciones con texto del archivo. Lanza FormatoNoSoportado o SinTexto."""
    extension = Path(nombre_archivo).suffix.lower()
    extractor = EXTRACTORES.get(extension)
    if extractor is None:
        raise FormatoNoSoportado(
            f"Formato no soportado: {extension or 'sin extension'}. "
            f"Usa {', '.join(FORMATOS)}."
        )
    secciones = [s for s in extractor(datos) if s.texto.strip()]
    if not secciones:
        raise SinTexto("No se pudo extraer texto del archivo.")
    return secciones
