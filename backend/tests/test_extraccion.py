"""Extraccion de texto de cada formato soportado."""

import io

import pytest
from docx import Document

from app.services.extraccion import FormatoNoSoportado, SinTexto, extraer


def pdf_minimo(paginas: list[str]) -> bytes:
    """PDF de una linea de texto por pagina, armado a mano (sin tabla xref: pypdf la reconstruye)."""
    objetos = ["<< /Type /Catalog /Pages 2 0 R >>"]
    ids_paginas = [4 + 2 * i for i in range(len(paginas))]
    kids = " ".join(f"{i} 0 R" for i in ids_paginas)
    objetos.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(paginas)} >>")
    objetos.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, texto in enumerate(paginas):
        flujo = f"BT /F1 12 Tf 20 100 Td ({texto}) Tj ET"
        objetos.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        )
        objetos.append(f"<< /Length {len(flujo)} >>\nstream\n{flujo}\nendstream")
    cuerpo = "%PDF-1.4\n"
    for n, obj in enumerate(objetos, 1):
        cuerpo += f"{n} 0 obj\n{obj}\nendobj\n"
    cuerpo += f"trailer\n<< /Root 1 0 R /Size {len(objetos) + 1} >>\nstartxref\n0\n%%EOF\n"
    return cuerpo.encode("latin-1")


def test_txt():
    secciones = extraer("nota.txt", b"Batman vive en Gotham.")
    assert len(secciones) == 1 and "Gotham" in secciones[0].texto


def test_txt_latin1():
    secciones = extraer("nota.txt", "Año de creación".encode("latin-1"))
    assert "Año" in secciones[0].texto


def test_md_se_parte_por_titulos():
    md = b"# Batman\nVive en Gotham.\n\n## Poderes\nNinguno.\n\n## Aliados\nRobin.\n"
    secciones = extraer("batman.md", md)
    assert [s.nombre for s in secciones if s.texto.strip()] == ["Batman", "Poderes", "Aliados"]
    assert secciones[1].texto.strip() == "Ninguno."


def test_csv_fila_a_fila():
    datos = b"nombre,ciudad\nBatman,Gotham\nSuperman,Metropolis\n"
    secciones = extraer("heroes.csv", datos)
    assert len(secciones) == 2
    assert secciones[0].texto == "nombre: Batman\nciudad: Gotham"


def test_csv_con_punto_y_coma():
    secciones = extraer("h.csv", b"nombre;ciudad\nBatman;Gotham\nFlash;Central\n")
    assert "ciudad: Central" in secciones[1].texto


def test_docx_por_titulos():
    doc = Document()
    doc.add_heading("Flash", level=1)
    doc.add_paragraph("Corre muy rapido.")
    doc.add_heading("Origen", level=2)
    doc.add_paragraph("Un rayo cayo sobre el.")
    buffer = io.BytesIO()
    doc.save(buffer)
    secciones = extraer("flash.docx", buffer.getvalue())
    nombres = [s.nombre for s in secciones]
    assert "Flash" in nombres and "Origen" in nombres
    assert any("rayo" in s.texto for s in secciones)


def test_pdf_por_pagina():
    secciones = extraer("dc.pdf", pdf_minimo(["Aquaman rules Atlantis", "Wonder Woman is Amazon"]))
    assert [s.nombre for s in secciones] == ["Pagina 1", "Pagina 2"]
    assert "Atlantis" in secciones[0].texto


def test_pdf_danado():
    with pytest.raises(SinTexto):
        extraer("roto.pdf", b"esto no es un pdf")


def test_formato_no_soportado():
    with pytest.raises(FormatoNoSoportado):
        extraer("video.mp4", b"\x00\x01")


def test_archivo_sin_texto():
    with pytest.raises(SinTexto):
        extraer("vacio.txt", b"   \n  ")
